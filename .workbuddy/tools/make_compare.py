# -*- coding: utf-8 -*-
"""用同一份真实大纲 + C 档字数契约生成对照 docx（WorkBuddy 版）"""
import json, io, re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SRC = r"C:/Users/zelin/WorkBuddy/2026-09-24-12-08-46/ab/result_ce_real.json"
OUT = r"C:/Users/zelin/Desktop/抖音商务衬衫polo选品分析-对照版-C档.docx"

rows = json.load(io.open(SRC, encoding="utf-8"))
rows = [r for r in rows if r["tier"] == "C" and not str(r["words"]).startswith("[")]

# 按原始顺序重建章节树
chapters, seen = [], {}
for r in rows:
    if r["chapter"] not in seen:
        seen[r["chapter"]] = {"title": r["chapter"], "paras": []}
        chapters.append(seen[r["chapter"]])
    seen[r["chapter"]]["paras"].append(r)


def set_cn(run, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


doc = Document()
st = doc.styles["Normal"]
st.font.name = "宋体"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

h = doc.add_heading("抖音商务衬衫polo选品分析", 0)
for r in h.runs:
    set_cn(r, "微软雅黑")

stats = []
for ci, ch in enumerate(chapters, 1):
    hh = doc.add_heading("%d %s" % (ci, ch["title"]), 1)
    for r in hh.runs:
        set_cn(r, "微软雅黑")
    for pi, r in enumerate(ch["paras"], 1):
        ph = doc.add_heading("%d.%d %s" % (ci, pi, r["title"]), 2)
        for run in ph.runs:
            set_cn(run, "微软雅黑")
        body = r["text"].strip()
        p = doc.add_paragraph(body)
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            set_cn(run, "宋体")
        dev = (r["words"] - r["target"]) / r["target"] * 100
        stats.append((r["title"], r["points"], r["target"], r["words"], dev))

doc.save(OUT)

w = [s[3] for s in stats]
hit = sum(1 for s in stats if abs(s[4]) <= 20)
tgt = [s[2] for s in stats]
print("OUT:", OUT)
print("段数 %d 均值 %.1f 中位 %d 最小 %d 最大 %d" % (
    len(w), sum(w) / len(w), sorted(w)[len(w) // 2], min(w), max(w)))
print("目标均值 %.1f  命中率(±20%%) %d/%d = %.0f%%" % (
    sum(tgt) / len(tgt), hit, len(stats), 100.0 * hit / len(stats)))
print("\n明细（段题 / 要点 / 目标 / 实际 / 偏差）")
for s in stats:
    print("  %-22s 点%d 目标%3d 实际%3d %+5.0f%%" % s)
