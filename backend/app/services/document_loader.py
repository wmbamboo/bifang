from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, Optional


ProgressCb = Optional[Callable[[int, int, str], None]]
ContinueCb = Optional[Callable[[], bool]]


def load_file_text(
    path: Path,
    *,
    enable_ocr: bool = False,
    on_progress: ProgressCb = None,
    should_continue: ContinueCb = None,
) -> str:
    """加载文件正文。enable_ocr=True 时对 PDF/图片走 OCR 增强（图文混排/扫描件）。"""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".log"}:
        return _read_text(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(
            path,
            enable_ocr=enable_ocr,
            on_progress=on_progress,
            should_continue=should_continue,
        )
    if suffix in {".pptx", ".ppt"}:
        return _read_pptx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
        if enable_ocr:
            from app.services.ocr_service import extract_image_text

            return extract_image_text(path)
        return ""
    return _read_text(path)


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except Exception:
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf_pypdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join([(p.extract_text() or "") for p in reader.pages])


def _read_pdf(
    path: Path,
    *,
    enable_ocr: bool = False,
    on_progress: ProgressCb = None,
    should_continue: ContinueCb = None,
) -> str:
    # 优先 PyMuPDF（文字层更稳）；需要 OCR 或 pymupdf 不可用时再分支
    if enable_ocr:
        from app.services.ocr_service import extract_pdf_text

        text, _meta = extract_pdf_text(
            path,
            enable_ocr=True,
            on_progress=on_progress,
            should_continue=should_continue,
        )
        return text

    try:
        import pymupdf as fitz

        doc = fitz.open(str(path))
        try:
            parts = []
            for i in range(doc.page_count):
                t = (doc.load_page(i).get_text("text") or "").strip()
                if t:
                    parts.append(f"[第{i + 1}页]\n{t}")
            return "\n\n".join(parts)
        finally:
            doc.close()
    except Exception:
        return _read_pdf_pypdf(path)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                t = shape.text.strip()
                if t:
                    slide_texts.append(t)
        if slide_texts:
            parts.append(f"[幻灯片{idx}]\n" + "\n".join(slide_texts))
    return "\n\n".join(parts)


def split_text(text: str, chunk_size: int = 500, chunk_overlap: int = 80) -> list[str]:
    return [c for c, _ in split_text_with_meta(text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)]


def _split_plain_text(
    body: str,
    *,
    page: Optional[int],
    chunk_size: int,
    chunk_overlap: int,
) -> list[tuple[str, dict]]:
    """普通正文按段/长度切；不含表。"""
    paragraphs = [p.strip() for p in body.split("\n") if p.strip()]
    out: list[tuple[str, dict]] = []
    buf = ""

    def _meta(kind: str = "text") -> dict:
        m: dict = {"kind": kind}
        if page is not None:
            m["page"] = page
        return m

    for p in paragraphs:
        if len(buf) + len(p) + 1 <= chunk_size:
            buf = f"{buf}\n{p}".strip() if buf else p
            continue
        if buf:
            out.append((buf, _meta()))
        if len(p) <= chunk_size:
            if out and chunk_overlap > 0:
                prev = out[-1][0]
                overlap = prev[-chunk_overlap:]
                buf = f"{overlap}\n{p}".strip()
            else:
                buf = p
        else:
            start = 0
            while start < len(p):
                end = start + chunk_size
                out.append((p[start:end], _meta()))
                start = max(end - chunk_overlap, end)
            buf = ""
    if buf:
        out.append((buf, _meta()))
    return out


def split_text_with_meta(
    text: str,
    chunk_size: int = 500,
    chunk_overlap: int = 80,
) -> list[tuple[str, dict]]:
    """切块并尽量继承页码。识别 `[第N页]` / `[幻灯片N]` 标记。

    表格整块入 chunk（kind=table），不按 token 数切开——根治表头/数据分家。
    """
    text = (text or "").strip()
    if not text:
        return []

    page_re = re.compile(r"^\[(?:第|幻灯片)?\s*(\d+)\s*页?\]\s*", re.M)
    parts: list[tuple[Optional[int], str]] = []
    cur_page: Optional[int] = None
    buf_lines: list[str] = []

    def _flush() -> None:
        nonlocal buf_lines, cur_page
        body = "\n".join(buf_lines).strip()
        if body:
            parts.append((cur_page, body))
        buf_lines = []

    for line in text.replace("\r\n", "\n").split("\n"):
        m = page_re.match(line.strip())
        if m:
            _flush()
            cur_page = int(m.group(1))
            rest = page_re.sub("", line.strip(), count=1).strip()
            if rest:
                buf_lines.append(rest)
            continue
        if line.strip():
            buf_lines.append(line.strip())
    _flush()

    if not parts:
        parts = [(None, text)]

    from app.services.table_structure import iter_page_segments

    out: list[tuple[str, dict]] = []
    for page, body in parts:
        for kind, seg, _s, _e in iter_page_segments(body):
            if kind == "table":
                meta: dict = {"kind": "table", "atomic": True}
                if page is not None:
                    meta["page"] = page
                out.append((seg, meta))
            else:
                out.extend(
                    _split_plain_text(
                        seg,
                        page=page,
                        chunk_size=chunk_size,
                        chunk_overlap=chunk_overlap,
                    )
                )
    return out
